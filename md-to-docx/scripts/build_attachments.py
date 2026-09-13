#!/usr/bin/env python3
"""Append an attachment section (证据册/附件册) to a pandoc-generated Chinese legal docx.

Usage:
    python3 build_attachments.py <docx> <manifest.json> [--align-signature]

manifest.json format:
{
  "title": "附 件",
  "attachments": [
    {"title": "附件1：申请人身份证复印件", "files": ["/path/img1.jpg"]},
    {"title": "附件2：xxx", "files": ["/path/a.jpg", "/path/b.jpg"]}
  ]
}

What it does:
- Forces A4 page size (21 x 29.7 cm, margins top/bottom 2.54 cm, left/right 3.18 cm)
- Appends a standalone attachment section: title page (H1 + file list) then one
  section per attachment (page break merged into the H2 title, keep_with_next)
- Inserts each image centered, scaled to max 5.7 in wide / 9.2 in high
- Sets image paragraphs to single line spacing + zero spacing (1.5x line
  spacing inflates inline image line height and pushes images off-page)
- Adds centered page footer "第 X 页" (宋体 10pt)
- Optional --align-signature: right-align paragraphs starting with
  "申请人：" or a Chinese date "YYYY年"

IMPORTANT: run AFTER scripts/format_docx.py — format_docx.py resets every
paragraph to 1.5x spacing, which would re-break image layouts.
"""
import argparse
import json
import sys

from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MAX_W_IN, MAX_H_IN = 5.7, 9.2   # A4 usable area: 5.76 x 9.69 in (margins 2.54/3.18 cm)


def set_style(p, style_id):
    pPr = p._p.get_or_add_pPr()
    el = pPr.makeelement(qn('w:pStyle'), {})
    el.set(qn('w:val'), style_id)
    pPr.insert(0, el)


def set_a4(section):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(3.18)


def add_page_footer(doc):
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.text = ''
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = footer_p.add_run('第 ')
    r2 = footer_p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    r2._r.append(f1); r2._r.append(instr); r2._r.append(f2)
    r3 = footer_p.add_run(' 页')
    for run in footer_p.runs:
        run.font.name = '宋体'
        run.font.size = Pt(10)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            rFonts.set(qn(attr), '宋体')


def align_signature(doc):
    import re
    date_only = re.compile(r'^\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日$')
    for para in doc.paragraphs:
        t = para.text.strip()
        if t.startswith('申请人：') or date_only.match(t):
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_picture_scaled(p_doc, path):
    img = Image.open(path)
    w, h = img.size
    w_in, h_in = w / 96.0, h / 96.0
    ratio = min(MAX_W_IN / w_in, MAX_H_IN / h_in, 1.0)
    w_in, h_in = w_in * ratio, h_in * ratio
    pp = p_doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(path, width=Inches(w_in), height=Inches(h_in))
    fmt = pp.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('manifest')
    ap.add_argument('--align-signature', action='store_true')
    args = ap.parse_args()

    with open(args.manifest, encoding='utf-8') as f:
        cfg = json.load(f)
    title = cfg.get('title', '附 件')
    attachments = cfg['attachments']

    doc = Document(args.docx)
    for s in doc.sections:
        set_a4(s)

    # Title page: H1 + file list (kept on one page)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph(title)
    set_style(h, 'Heading1')
    h.paragraph_format.keep_with_next = True
    lp = doc.add_paragraph()
    for att in attachments:
        lp.add_run(f"{att['title']}\n")

    # Per attachment: page break merged into title paragraph
    for att in attachments:
        h2 = doc.add_paragraph()
        set_style(h2, 'Heading2')
        r = h2.add_run()
        r.add_break(WD_BREAK.PAGE)
        h2.add_run(att['title'])
        h2.paragraph_format.keep_with_next = True
        for f in att['files']:
            add_picture_scaled(doc, f)

    if args.align_signature:
        align_signature(doc)
    add_page_footer(doc)

    doc.save(args.docx)
    print(f'Attachments appended: {args.docx}')


if __name__ == '__main__':
    sys.exit(main())
