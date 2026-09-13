#!/usr/bin/env python3
"""Post-process a pandoc-generated docx for Chinese legal document formatting.

Applies:
- 宋体 font, 12pt, black, for every run (body, headings, table cells)
- Main title (first Heading 1): centered, 16pt bold
- Other headings: bold, 12pt black (no theme blue)
- 1.5x line spacing, zero paragraph spacing everywhere
- Tables: full page width (100%), header row bold
"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH

TITLE_SIZE = Pt(16)
TITLE_STYLE = 'Heading 1'


def _set_run_font(run, size=Pt(12), bold=None, color=(0, 0, 0)):
    run.font.name = '宋体'
    run.font.size = size
    run.font.color.rgb = RGBColor(*color)
    if bold is not None:
        run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), '宋体')


def _set_para_format(para, bold=None, size=Pt(12), align=None):
    fmt = para.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if align is not None:
        para.alignment = align
    for run in para.runs:
        _set_run_font(run, size=size, bold=bold)


def _set_normal_style(doc):
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), '宋体')


def _set_table_full_width(table):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = tblPr.makeelement(qn('w:tblW'), {})
        tblPr.append(tblW)
    tblW.set(qn('w:type'), 'pct')
    tblW.set(qn('w:w'), '5000')


def _add_table_borders(table):
    """Explicit 0.5pt single borders on all edges — Word does not show
    borders unless they are declared (LibreOffice does, hence the
    LibreOffice-OK / Word-invisible discrepancy)."""
    tblPr = table._tbl.tblPr
    if tblPr.find(qn('w:tblBorders')) is not None:
        return
    tblBorders = tblPr.makeelement(qn('w:tblBorders'), {})
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = tblPr.makeelement(qn('w:' + edge), {})
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        tblBorders.append(el)
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is not None:
        tblW.addnext(tblBorders)
    else:
        tblPr.append(tblBorders)


def _set_first_line_indent(doc):
    """Body paragraphs (pandoc 'Body Text'/'First Paragraph') get a 2-char
    first-line indent, the norm for Chinese formal documents. Excludes:
    address line ("致：…"), signature block (right-aligned), list items
    ('Compact'), image paragraphs, empty paragraphs."""
    from docx.oxml.ns import qn as _qn
    for para in doc.paragraphs:
        if para.style.name not in ('Body Text', 'First Paragraph'):
            continue
        t = para.text.strip()
        if not t or t.startswith('致：'):
            continue
        if para.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            continue
        if para._p.findall('.//' + _qn('w:drawing')):
            continue
        para.paragraph_format.first_line_indent = Pt(24)


def process_docx(path: str) -> None:
    doc = Document(path)

    title_done = False
    for para in doc.paragraphs:
        if para.style.name == TITLE_STYLE:
            if not title_done:
                _set_para_format(para, bold=True, size=TITLE_SIZE,
                                 align=WD_ALIGN_PARAGRAPH.CENTER)
                title_done = True
            else:
                _set_para_format(para, bold=True)
        elif para.style.name.startswith('Heading'):
            _set_para_format(para, bold=True)
        else:
            _set_para_format(para)

    for table in doc.tables:
        _set_table_full_width(table)
        _add_table_borders(table)
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                for para in cell.paragraphs:
                    _set_para_format(para, bold=True if ri == 0 else None)

    _set_normal_style(doc)
    _set_first_line_indent(doc)
    doc.save(path)
    print(f'Formatted: {path}')


if __name__ == '__main__':
    process_docx(sys.argv[1])
